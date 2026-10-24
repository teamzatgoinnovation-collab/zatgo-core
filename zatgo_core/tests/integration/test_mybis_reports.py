"""MyBis reports, exports, global search and the notification diff against real documents."""

from __future__ import annotations

import base64
import csv
import io
from unittest.mock import patch

import frappe
from frappe.tests.classes.integration_test_case import IntegrationTestCase
from frappe.utils import add_days, getdate, random_string, today

from zatgo_core.api.v1.mybis import alerts as alerts_api
from zatgo_core.api.v1.mybis import reports as rep_api
from zatgo_core.api.v1.mybis import search as search_api
from zatgo_core.services.mybis import notifications, reports
from zatgo_core.tests.integration.test_mybis_analytics import _purchase, _supplier
from zatgo_core.tests.integration.test_mybis_dashboard import COMPANY, _company, _customer, _group, _invoice, _item


def _decode(env):
    return base64.b64decode(env["data"]["content_base64"])


class TestMyBisReports(IntegrationTestCase):
    @classmethod
    def setUpClass(cls) -> None:
        super().setUpClass()
        frappe.set_user("Administrator")
        _company()
        frappe.db.commit()

    def setUp(self) -> None:
        frappe.set_user("Administrator")
        self.cat = _group("MyBis Cat A")
        self.item = _item(self.cat)
        self.cust, self.cust2 = _customer(), _customer()
        self.sup = _supplier()
        self.first = getdate(today()).replace(day=1)
        self.inv1 = _invoice(self.cust, self.item, 2, 100, self.first)  # 200
        self.inv2 = _invoice(self.cust2, self.item, 1, 50, self.first)  # 50
        self.inv_old = _invoice(self.cust, self.item, 1, 100, add_days(self.first, -3))  # previous period
        self.pi = _purchase(self.sup, self.item, 4, 25, self.first)

    def tearDown(self) -> None:
        frappe.set_user("Administrator")
        frappe.db.rollback()

    def run_report(self, report, **kw):
        return rep_api.run(report=report, company=COMPANY, preset="this_month", refresh=1, **kw)

    # -- catalog -----------------------------------------------------------------------------------------------------------

    def test_catalog_covers_the_requested_reports_and_marks_availability(self) -> None:
        d = rep_api.catalog()["data"]
        keys = {r["key"] for r in d["reports"]}
        self.assertEqual(
            keys,
            {"sales_summary", "sales_by_customer", "sales_by_item", "sales_by_category", "salesperson_performance", "purchase_summary", "purchase_by_supplier", "purchase_by_item",
             "stock_summary", "stock_movement", "stock_aging", "slow_moving", "fast_moving", "profit_and_loss", "balance_sheet", "cash_flow", "accounts_receivable", "accounts_payable", "general_ledger"},
        )
        self.assertEqual({r["group"] for r in d["reports"]}, {"Sales", "Purchases", "Inventory", "Finance"})
        self.assertTrue(all(r["available"] for r in d["reports"]))
        self.assertEqual(d["formats"], ["csv", "xlsx", "pdf"])
        self.assertEqual({r["key"]: r["source"] for r in d["reports"]}["profit_and_loss"], "erpnext")

    def test_a_report_the_caller_may_not_use_is_marked_unavailable_and_refused(self) -> None:
        with patch("frappe.core.doctype.report.report.Report.is_permitted", return_value=False):
            d = rep_api.catalog()["data"]
        self.assertFalse(next(r for r in d["reports"] if r["key"] == "profit_and_loss")["available"])
        self.assertTrue(next(r for r in d["reports"] if r["key"] == "sales_summary")["available"])
        with patch("frappe.core.doctype.report.report.Report.is_permitted", return_value=False):
            with self.assertRaises(frappe.PermissionError):
                self.run_report("general_ledger")

    # -- MyBis aggregate reports -------------------------------------------------------------------------------------------

    def test_sales_by_customer_has_typed_columns_totals_growth_and_notes(self) -> None:
        r = self.run_report("sales_by_customer")["data"]
        self.assertEqual({c["key"]: c["type"] for c in r["columns"]}["revenue"], "currency")
        rows = {x["code"]: x for x in r["rows"]}
        self.assertEqual((rows[self.cust]["revenue"], rows[self.cust]["previous_revenue"], rows[self.cust]["growth_percent"]), (200.0, 100.0, 100.0))
        self.assertEqual(rows[self.cust2]["share_percent"], 20.0)
        self.assertEqual(r["totals"], {"name": "Total", "revenue": 250.0})
        self.assertEqual((r["currency"], r["company"], r["total_rows"], r["source"]), ("SAR", COMPANY, 2, "mybis"))
        self.assertTrue(r["notes"])

    def test_summary_and_other_aggregates(self) -> None:
        s = self.run_report("sales_summary")["data"]
        by = {x["metric"]: x for x in s["rows"]}
        self.assertEqual((by["Gross sales"]["value"], by["Invoices"]["value"]), (250.0, 2))
        p = self.run_report("purchase_by_supplier")["data"]
        self.assertEqual((p["rows"][0]["code"], p["rows"][0]["revenue"]), (self.sup, 100.0))
        self.assertNotIn("profit", p["rows"][0])
        self.assertEqual(self.run_report("sales_by_category")["data"]["rows"][0]["code"], self.cat)
        self.assertEqual(self.run_report("salesperson_performance")["data"]["rows"], [])

    # -- ERPNext-backed reports --------------------------------------------------------------------------------------------

    def test_profit_and_loss_is_erpnexts_with_indented_rows(self) -> None:
        r = self.run_report("profit_and_loss")["data"]
        self.assertEqual(r["source_report"], "Profit and Loss Statement")
        total = next(x for x in r["rows"] if x["label"] == "Total Income (Credit)")
        self.assertEqual(total["value"], 250.0)
        self.assertTrue(any(x["_indent"] >= 1 for x in r["rows"]))

    def test_general_ledger_receivable_and_stock_reports_normalise_erpnext_output(self) -> None:
        gl = self.run_report("general_ledger")["data"]
        keys = {c["key"]: c["type"] for c in gl["columns"]}
        self.assertEqual((keys["debit"], keys["credit"], keys["posting_date"]), ("currency", "currency", "date"))
        self.assertNotIn("gl_entry", keys)  # hidden columns are dropped
        self.assertTrue(any(x.get("account") == "Opening" for x in gl["rows"]))  # ERPNext's quoting is removed
        ar = self.run_report("accounts_receivable", filters={"party": self.cust})["data"]
        self.assertEqual({x["voucher_no"] for x in ar["rows"] if x.get("voucher_no")}, {self.inv1, self.inv_old})
        self.assertTrue(all(x.get("party") == self.cust for x in ar["rows"] if x.get("party")))
        self.assertTrue(self.run_report("accounts_payable")["data"]["rows"])
        for key in ("stock_summary", "stock_movement", "stock_aging", "balance_sheet", "cash_flow", "slow_moving", "fast_moving", "purchase_summary", "purchase_by_item", "sales_by_item"):
            t = self.run_report(key)["data"]
            self.assertIn("columns", t, key)
            self.assertIsInstance(t["rows"], list, key)

    def test_stock_summary_shows_stock_and_accepts_its_filters(self) -> None:
        abbr = frappe.db.get_value("Company", COMPANY, "abbr")
        wh = f"MyBis Stock - {abbr}"
        if not frappe.db.exists("Warehouse", wh):
            frappe.get_doc({"doctype": "Warehouse", "warehouse_name": "MyBis Stock", "company": COMPANY}).insert(ignore_permissions=True)
        code = f"MYBIS-RPT-{random_string(6).upper()}"
        frappe.get_doc({"doctype": "Item", "item_code": code, "item_name": code, "item_group": self.cat, "stock_uom": "Nos", "is_stock_item": 1}).insert(ignore_permissions=True)
        frappe.get_doc({"doctype": "Stock Entry", "stock_entry_type": "Material Receipt", "company": COMPANY, "posting_date": self.first, "set_posting_time": 1, "items": [{"item_code": code, "qty": 7, "basic_rate": 10, "t_warehouse": wh}]}).insert(ignore_permissions=True).submit()
        t = self.run_report("stock_summary", filters={"item_code": code})["data"]
        row = next(x for x in t["rows"] if x.get("item_code") == code)
        self.assertEqual((row["bal_qty"], row["in_qty"]), (7.0, 7.0))
        mv = self.run_report("stock_movement", filters={"item_code": code})["data"]
        self.assertTrue(any(x.get("item_code") == code and x.get("in_qty") == 7.0 for x in mv["rows"]))
        self.assertIsInstance(self.run_report("stock_aging", filters={"item_code": code})["data"]["rows"], list)

    # -- filters / pagination ----------------------------------------------------------------------------------------------

    def test_pagination_and_input_checks(self) -> None:
        p = self.run_report("general_ledger", page_size=3)["data"]
        self.assertEqual((len(p["rows"]), p["page"], p["page_size"]), (3, 1, 3))
        self.assertGreater(p["total_rows"], 3)
        p2 = self.run_report("general_ledger", page=2, page_size=3)["data"]
        self.assertNotEqual(p["rows"], p2["rows"])
        with self.assertRaises(frappe.ValidationError):
            self.run_report("nope")
        with self.assertRaises(frappe.ValidationError):
            self.run_report("sales_summary", filters={"party": "x"})  # not an allowed filter of this report
        with self.assertRaises(frappe.ValidationError):
            self.run_report("general_ledger", filters="not json")

    def test_the_result_is_cached_with_its_age(self) -> None:
        a = rep_api.run(report="sales_summary", company=COMPANY, refresh=1)
        b = rep_api.run(report="sales_summary", company=COMPANY)
        self.assertFalse(a["meta"]["cached"])
        self.assertTrue(b["meta"]["cached"])

    # -- exports -----------------------------------------------------------------------------------------------------------

    def test_csv_export_matches_the_table_and_opens_in_excel(self) -> None:
        env = rep_api.export(report="sales_by_customer", format="csv", company=COMPANY, preset="this_month")
        raw = _decode(env)
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))  # BOM for Excel
        self.assertTrue(env["data"]["filename"].endswith(".csv"))
        self.assertIn(COMPANY.replace(" ", "_"), env["data"]["filename"])
        rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
        self.assertEqual(rows[0], ["Sales by Customer"])
        header = next(r for r in rows if r and r[0] == "Customer")
        data = [r for r in rows if r and r[0] and r[0] not in ("Customer", "Sales by Customer", "Total") and len(r) > 3]
        self.assertEqual(header[:3], ["Customer", "Code", "Net sales"])
        self.assertEqual(sum(float(r[2]) for r in data if r[1] in (self.cust, self.cust2)), 250.0)
        self.assertEqual(next(r for r in rows if r and r[0] == "Total")[2], "250.0")

    def test_xlsx_export_is_a_real_workbook_with_numbers_as_numbers(self) -> None:
        from openpyxl import load_workbook

        env = rep_api.export(report="sales_by_customer", format="xlsx", company=COMPANY, preset="this_month")
        wb = load_workbook(io.BytesIO(_decode(env)))
        ws = wb.active
        self.assertEqual(ws.title, "Sales by Customer")
        cells = [[c.value for c in row] for row in ws.iter_rows()]
        header_row = next(i for i, r in enumerate(cells) if r and r[0] == "Customer")
        mine = next(r for r in cells[header_row + 1 :] if r[2] == 200)
        self.assertEqual(mine[1], self.cust)
        self.assertEqual(mine[2], 200.0)
        self.assertIsInstance(mine[2], (int, float))

    def test_pdf_export_is_a_pdf_and_refuses_what_would_not_fit(self) -> None:
        env = rep_api.export(report="profit_and_loss", format="pdf", company=COMPANY, preset="this_month")
        raw = _decode(env)
        self.assertTrue(raw.startswith(b"%PDF"))
        self.assertEqual(env["data"]["size"], len(raw))
        with patch.object(reports, "PDF_MAX_ROWS", 1):
            with self.assertRaises(frappe.ValidationError):
                rep_api.export(report="general_ledger", format="pdf", company=COMPANY, preset="this_month")

    def test_exports_neutralise_spreadsheet_formulas(self) -> None:
        evil = frappe.get_doc({"doctype": "Customer", "customer_name": "=HYPERLINK(\"http://x\",\"pay\")", "customer_type": "Individual",
                               "customer_group": frappe.db.get_value("Customer Group", {"is_group": 0}, "name"), "territory": frappe.db.get_value("Territory", {"is_group": 0}, "name")}).insert(ignore_permissions=True)
        _invoice(evil.name, self.item, 1, 10, self.first)
        raw = _decode(rep_api.export(report="sales_by_customer", format="csv", company=COMPANY, preset="this_month")).decode("utf-8-sig")
        self.assertIn("'=HYPERLINK", raw)
        self.assertNotIn("\n=HYPERLINK", raw)
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(_decode(rep_api.export(report="sales_by_customer", format="xlsx", company=COMPANY, preset="this_month")))).active
        values = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        self.assertTrue(any(v.startswith("'=HYPERLINK") for v in values))
        self.assertFalse(any(v.startswith("=") for v in values))

    def test_export_input_is_checked(self) -> None:
        with self.assertRaises(frappe.ValidationError):
            rep_api.export(report="sales_summary", format="docx", company=COMPANY)
        with patch.object(reports, "EXPORT_MAX_BYTES", 10):
            with self.assertRaises(frappe.ValidationError):
                rep_api.export(report="sales_summary", format="csv", company=COMPANY)

    def test_unit_level_helpers(self) -> None:
        cols = reports.normalise_columns(["Item Code:Link/Item:120", "Qty:Float:80", {"label": "Hidden", "fieldname": "h", "hidden": 1}, {"label": "Amount", "fieldname": "amt", "fieldtype": "Currency"}])
        self.assertEqual([(c["key"], c["type"]) for c in cols], [("item_code", "text"), ("qty", "number"), ("amt", "currency")])
        t = reports.normalise(["Name:Data", "Total:Currency"], [["a", 1.5], ["'Total'", 2]])
        self.assertEqual(t["rows"], [{"name": "a", "total": 1.5}, {"name": "Total", "total": 2.0}])
        self.assertEqual(reports.safe_cell("=1+1"), "'=1+1")
        self.assertEqual(reports.safe_cell(" -x"), "' -x")
        self.assertEqual(reports.safe_cell("Acme"), "Acme")
        self.assertEqual(reports.safe_cell(-5.0), -5.0)


class TestMyBisSearch(IntegrationTestCase):
    @classmethod
    def setUpClass(cls) -> None:
        super().setUpClass()
        frappe.set_user("Administrator")
        _company()
        frappe.db.commit()

    def setUp(self) -> None:
        frappe.set_user("Administrator")
        self.tag = random_string(6).lower()
        self.cust = frappe.get_doc({"doctype": "Customer", "customer_name": f"Zzq {self.tag} Traders", "customer_type": "Company",
                                    "customer_group": frappe.db.get_value("Customer Group", {"is_group": 0}, "name"), "territory": frappe.db.get_value("Territory", {"is_group": 0}, "name")}).insert(ignore_permissions=True).name
        self.sup = frappe.get_doc({"doctype": "Supplier", "supplier_name": f"Zzq {self.tag} Supplies", "supplier_group": frappe.db.get_value("Supplier Group", {"is_group": 0}, "name")}).insert(ignore_permissions=True).name
        self.item = _item(_group("MyBis Cat A"))
        frappe.db.set_value("Item", self.item, "item_name", f"Zzq {self.tag} Widget")
        self.inv = _invoice(self.cust, self.item, 1, 40, getdate(today()).replace(day=1))

    def tearDown(self) -> None:
        frappe.set_user("Administrator")
        frappe.db.rollback()

    def q(self, text, **kw):
        return search_api.query(q=text, **kw)["data"]

    def test_finds_each_record_type_and_gives_each_a_drill_target(self) -> None:
        d = self.q(f"zzq {self.tag}")
        by = {h["type"]: h for h in d["results"]}
        self.assertEqual(by["customer"]["id"], self.cust)
        self.assertEqual(by["customer"]["drill"], {"screen": "customers", "dimension": "customer", "filters": {"customer": self.cust}, "title": f"Zzq {self.tag} Traders"})
        self.assertEqual(by["supplier"]["drill"]["filters"], {"supplier": self.sup})
        self.assertEqual(by["item"]["drill"]["filters"], {"product": self.item})
        inv = self.q(self.inv, types=["sales_invoice"])["results"][0]
        self.assertEqual((inv["id"], inv["amount"], inv["drill"]), (self.inv, 40.0, {"screen": "sales", "invoice": self.inv}))
        self.assertEqual(inv["subtitle"], f"Zzq {self.tag} Traders")
        self.assertEqual({h["type"] for h in self.q(self.cust, types="sales_invoice,customer")["results"]}, {"customer", "sales_invoice"})

    def test_matches_on_the_party_name_of_an_invoice_and_ranks_exact_first(self) -> None:
        hits = self.q(f"Zzq {self.tag} Traders", types=["sales_invoice", "customer"])["results"]
        self.assertEqual({h["type"] for h in hits}, {"customer", "sales_invoice"})
        self.assertEqual(hits[0]["type"], "customer")  # an exact title match before the invoice that merely mentions it

    def test_company_filter_limits_company_records(self) -> None:
        mine = self.q(self.inv, types=["sales_invoice"], company=COMPANY)["results"]
        self.assertEqual(len(mine), 1)
        self.assertEqual(self.q(self.inv, types=["sales_invoice"], company="ZatGo Core Test Co")["results"], [])

    def test_short_queries_unknown_types_and_limits(self) -> None:
        with self.assertRaises(frappe.ValidationError):
            self.q("a")
        with self.assertRaises(frappe.ValidationError):
            self.q("zzq", types=["salaries"])
        for i in range(3):
            frappe.get_doc({"doctype": "Customer", "customer_name": f"Zzq {self.tag} Extra {i}", "customer_type": "Individual", "customer_group": frappe.db.get_value("Customer Group", {"is_group": 0}, "name"), "territory": frappe.db.get_value("Territory", {"is_group": 0}, "name")}).insert(ignore_permissions=True)
        self.assertEqual(len(self.q(f"zzq {self.tag}", types=["customer"], limit=2)["results"]), 2)

    def test_wildcards_and_injection_are_only_text(self) -> None:
        self.assertEqual(self.q("x' OR '1'='1", types=["customer"])["results"], [])
        self.assertIsInstance(self.q("%%", types=["customer"])["results"], list)

    def test_unreadable_types_are_skipped_not_searched_around(self) -> None:
        real = frappe.has_permission
        with patch("zatgo_core.services.mybis.search.frappe.has_permission", side_effect=lambda dt, *a, **k: dt != "Supplier" and real(dt, *a, **k)):
            d = self.q(f"zzq {self.tag}")
        self.assertIn("supplier", d["skipped"])
        self.assertNotIn("supplier", {h["type"] for h in d["results"]})
        self.assertIn("customer", {h["type"] for h in d["results"]})

    def test_a_company_restricted_user_never_sees_another_companys_invoices(self) -> None:
        from zatgo_core.tests.integration._fixtures import get_or_create_test_company

        other = get_or_create_test_company()
        email = f"mybis.srch.{random_string(8).lower()}@zatgo.test"
        frappe.get_doc({"doctype": "User", "email": email, "first_name": "S", "send_welcome_email": 0, "roles": [{"role": "System Manager"}, {"role": "Accounts Manager"}]}).insert(ignore_permissions=True)
        frappe.get_doc({"doctype": "User Permission", "user": email, "allow": "Company", "for_value": other}).insert(ignore_permissions=True)
        frappe.clear_cache(user=email)
        frappe.set_user(email)
        self.assertEqual(self.q(self.inv, types=["sales_invoice"])["results"], [])
        with self.assertRaises(frappe.PermissionError):
            self.q(self.inv, types=["sales_invoice"], company=COMPANY)

    def test_record_summaries(self) -> None:
        abbr = frappe.db.get_value("Company", COMPANY, "abbr")
        wh = f"MyBis Rec - {abbr}"
        if not frappe.db.exists("Warehouse", wh):
            frappe.get_doc({"doctype": "Warehouse", "warehouse_name": "MyBis Rec", "company": COMPANY}).insert(ignore_permissions=True)
        r = search_api.record(type="warehouse", name=wh)["data"]
        self.assertEqual((r["label"], r["currency"]), ("Warehouse", "SAR"))
        self.assertTrue({"Stock quantity", "Stock value"} <= {f["label"] for f in r["fields"]})
        with self.assertRaises(frappe.ValidationError):
            search_api.record(type="customer", name=self.cust)  # customers have their own drill-down screen
        with self.assertRaises(frappe.DoesNotExistError):
            search_api.record(type="warehouse", name="nope")


class TestMyBisNotifications(IntegrationTestCase):
    ITEMS = [
        {"id": "a", "severity": "critical", "fingerprint": "a|co|critical|2026-10-01"},
        {"id": "b", "severity": "medium", "fingerprint": "b|co|medium|2026-10-01"},
        {"id": "c", "severity": "low", "fingerprint": "c|co|low|2026-10-01"},
    ]

    def test_unread_is_new_or_escalated_and_respects_the_minimum_severity(self) -> None:
        ids = lambda items: [i["id"] for i in items]  # noqa: E731
        self.assertEqual(ids(notifications.unread(self.ITEMS, [])), ["a", "b", "c"])
        self.assertEqual(ids(notifications.unread(self.ITEMS, ["b|co|medium|2026-10-01"])), ["a", "c"])
        self.assertEqual(ids(notifications.unread(self.ITEMS, [], "medium")), ["a", "b"])
        self.assertEqual(ids(notifications.unread([{**self.ITEMS[1], "sid": "abc123"}], ["abc123"])), [])  # the short id works too
        escalated = [{"id": "b", "severity": "high", "fingerprint": "b|co|high|2026-10-01"}]  # the same condition got worse
        self.assertEqual(ids(notifications.unread(escalated, ["b|co|medium|2026-10-01"])), ["b"])

    def test_dispatch_goes_through_the_registered_channels(self) -> None:
        class Capture(notifications.NotificationChannel):
            name = "capture"
            got: list = []

            def deliver(self, user, items):
                Capture.got = [i["id"] for i in items]
                return len(items)

        with patch.dict(notifications.CHANNELS, {"capture": Capture()}):
            out = notifications.dispatch("u@x", self.ITEMS, ["c|co|low|2026-10-01"], "medium")
            self.assertEqual(out, {"in_app": 2, "capture": 2})
            self.assertEqual(Capture.got, ["a", "b"])
            self.assertEqual(notifications.dispatch("u@x", self.ITEMS, [i["fingerprint"] for i in self.ITEMS]), {"in_app": 0, "capture": 0})

    def test_the_feed_reports_unread_for_the_fingerprints_the_app_has_seen(self) -> None:
        _company()
        cat = _group("MyBis Cat A")
        item, cust = _item(cat), _customer()
        first = getdate(today()).replace(day=1)
        prev = getdate(add_days(first, -1)).replace(day=1)
        _invoice(cust, item, 10, 100, prev)
        _invoice(cust, item, 1, 100, first)
        try:
            full = alerts_api.feed(company=COMPANY, refresh=1)["data"]
            self.assertNotIn("unread", full)
            first_pass = alerts_api.feed(company=COMPANY, seen="[]")["data"]
            self.assertEqual(set(first_pass["unread"]), {i["id"] for i in first_pass["items"]})
            self.assertTrue(all(len(i["sid"]) == 10 for i in first_pass["items"]))
            seen = [i["sid"] for i in first_pass["items"]]
            again = alerts_api.feed(company=COMPANY, seen=frappe.as_json(seen))["data"]
            self.assertEqual(again["unread"], [])
            only_critical = alerts_api.feed(company=COMPANY, seen="[]", min_severity="critical")["data"]
            self.assertTrue(all(next(i for i in only_critical["items"] if i["id"] == u)["severity"] == "critical" for u in only_critical["unread"]))
            with self.assertRaises(frappe.ValidationError):
                alerts_api.feed(company=COMPANY, seen="{not a list}")
        finally:
            frappe.db.rollback()
