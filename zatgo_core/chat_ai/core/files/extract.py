"""Text extraction from uploaded documents (bytes in, text out)."""

from __future__ import annotations

import csv
import io
import os
import zipfile
from dataclasses import dataclass

MAX_CHARS = 1_000_000
MAX_PDF_PAGES = 300
MAX_SHEET_ROWS = 20000
MAX_ZIP_UNCOMPRESSED = 200 * 1024 * 1024  # decompression-bomb guard for docx/xlsx
MAX_ZIP_RATIO = 200


class ExtractError(ValueError):
	"""User-facing extraction failure."""


@dataclass
class Extracted:
	text: str
	pages: int = 0
	truncated: bool = False


def _decode(content: bytes) -> str:
	"""UTF-8 first; Excel on Arabic Windows exports CSV as Windows-1256."""
	for enc in ("utf-8-sig", "cp1256"):
		try:
			return content.decode(enc)
		except UnicodeDecodeError:
			continue
	return content.decode("latin-1")


def _check_zip(content: bytes) -> None:
	try:
		with zipfile.ZipFile(io.BytesIO(content)) as z:
			total = sum(i.file_size for i in z.infolist())
	except zipfile.BadZipFile as exc:
		raise ExtractError("The file is corrupt or not a valid Office document.") from exc
	if total > MAX_ZIP_UNCOMPRESSED or total > max(len(content), 1) * MAX_ZIP_RATIO:
		raise ExtractError("The file expands to an unsafe size and was rejected.")


def _cap(text: str) -> Extracted:
	if len(text) > MAX_CHARS:
		return Extracted(text[:MAX_CHARS], truncated=True)
	return Extracted(text)


def _opens_without_password(reader) -> bool:
	"""Many PDFs are "encrypted" only to restrict printing or copying (owner password) and
	open for any reader with an empty user password; only real user passwords block us."""
	try:
		return bool(reader.decrypt(""))
	except Exception:  # wrong password, or an unsupported / unavailable cipher
		return False


def extract_pdf(content: bytes) -> Extracted:
	from pypdf import PdfReader
	from pypdf.errors import PyPdfError

	try:
		reader = PdfReader(io.BytesIO(content))
		if reader.is_encrypted and not _opens_without_password(reader):
			raise ExtractError("The PDF is password protected.")
		parts, total = [], len(reader.pages)
		for i, page in enumerate(reader.pages[:MAX_PDF_PAGES], 1):
			text = (page.extract_text() or "").strip()
			if text:
				parts.append(f"[Page {i}]\n{text}")
	except ExtractError:
		raise
	except (PyPdfError, ValueError, KeyError, OSError) as exc:
		raise ExtractError("The PDF could not be read.") from exc
	if not parts:
		raise ExtractError(
			"This PDF has no extractable text (it may be a scan). Upload it as an image instead."
		)
	out = _cap("\n\n".join(parts))
	out.pages = total
	out.truncated = out.truncated or total > MAX_PDF_PAGES
	return out


def extract_docx(content: bytes) -> Extracted:
	import docx

	_check_zip(content)
	try:
		document = docx.Document(io.BytesIO(content))
	except Exception as exc:  # python-docx raises assorted errors on bad packages
		raise ExtractError("The document could not be read.") from exc
	parts = [p.text for p in document.paragraphs if p.text.strip()]
	for table in document.tables:
		for row in table.rows:
			cells = [c.text.strip().replace("\n", " ") for c in row.cells]
			if any(cells):
				parts.append(" | ".join(cells))
	if not parts:
		raise ExtractError("The document contains no text.")
	return _cap("\n".join(parts))


def read_xlsx_rows(content: bytes, *, max_rows: int = MAX_SHEET_ROWS) -> list[tuple[str, list[list]]]:
	"""[(sheet name, rows)] with computed values (formulas are not evaluated by us)."""
	import openpyxl

	_check_zip(content)
	try:
		wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
	except Exception as exc:
		raise ExtractError("The spreadsheet could not be read.") from exc
	sheets = []
	try:
		for ws in wb.worksheets:
			rows = []
			for i, row in enumerate(ws.iter_rows(values_only=True)):
				if i >= max_rows:
					break
				rows.append(["" if v is None else v for v in row])
			while rows and not any(str(v).strip() for v in rows[-1]):
				rows.pop()
			sheets.append((ws.title, rows))
	finally:
		wb.close()
	return sheets


def _rows_to_text(rows: list[list]) -> str:
	buf = io.StringIO()
	w = csv.writer(buf, lineterminator="\n")
	for r in rows:
		w.writerow(r)
	return buf.getvalue()


def extract_xlsx(content: bytes) -> Extracted:
	sheets = read_xlsx_rows(content)
	parts = [f"[Sheet: {name}]\n{_rows_to_text(rows)}" for name, rows in sheets if rows]
	if not parts:
		raise ExtractError("The spreadsheet is empty.")
	out = _cap("\n".join(parts))
	out.pages = len(sheets)
	return out


def extract_text(filename: str, content: bytes) -> Extracted:
	ext = os.path.splitext(filename)[1].lower()
	if ext == ".pdf":
		return extract_pdf(content)
	if ext == ".docx":
		return extract_docx(content)
	if ext == ".xlsx":
		return extract_xlsx(content)
	if ext in (".csv", ".txt", ".md"):
		text = _decode(content)
		if not text.strip():
			raise ExtractError("The file is empty.")
		return _cap(text)
	raise ExtractError("This file type has no text to extract.")


def parse_table(filename: str, content: bytes, *, max_rows: int = MAX_SHEET_ROWS) -> tuple[list[str], list[list]]:
	"""(columns, rows) of a CSV or the first non-empty sheet of an XLSX."""
	ext = os.path.splitext(filename)[1].lower()
	if ext == ".xlsx":
		rows = next((r for _n, r in read_xlsx_rows(content, max_rows=max_rows + 1) if r), [])
	elif ext == ".csv":
		text = _decode(content)
		try:
			dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
		except csv.Error:
			dialect = csv.excel
		rows = [r for _i, r in zip(range(max_rows + 1), csv.reader(io.StringIO(text), dialect))]
	else:
		raise ExtractError("Only CSV and XLSX files can be analysed as tables.")
	if not rows:
		raise ExtractError("The table is empty.")
	header = [str(h).strip() or f"column_{i + 1}" for i, h in enumerate(rows[0])]
	return header, [list(r) + [""] * (len(header) - len(r)) for r in rows[1 : max_rows + 1]]
